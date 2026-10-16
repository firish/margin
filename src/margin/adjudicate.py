"""Blinded adjudication of coder-vs-human disagreements (RQ1 corrected triage, RQ5 error analysis).

Items are *pooled* across coders: a unit enters the pool if any coder disagrees with the human
reference, so the corrected labels are fair to every coder rather than only the one adjudicated.
All high-confidence disagreements are taken first (they decide coverage at strict error targets),
then a random sample of the rest, plus hidden controls where every coder agrees with the human
(their overturn rate estimates how often the adjudicator overrules a label that is probably fine).

The adjudicator sees: segment, context, the codebook, and shuffled candidate labels. Never: which
candidate is the human label, coder names, or confidences. The answer key is written separately.

Units: exclusive codebooks -> one item per segment ("which candidate is correct?");
multi codebooks -> one item per (segment, code) ("does this code apply? Y/N/?").
"""

from __future__ import annotations

import json
import re
import string
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd

from margin.datasets.base import Task
from margin.metrics.core import reference
from margin.run import RUNS, Experiment
from margin.schema import Codebook

ADJUDICATION = Path("data/adjudication")
LETTERS = string.ascii_uppercase
NONE_ANSWERS = {"none", "neither", "x"}
MULTI_ANSWERS = {"y": 1, "yes": 1, "n": 0, "no": 0, "?": None}


@dataclass(frozen=True)
class ExportConfig:
    sources: tuple[tuple[str, str, str], ...]
    """(name, experiment, file stem) per coder, e.g. ("jev", "jev_main", "jev__canonical")."""
    tasks: tuple[str, ...]
    split: str | None = "test"
    n_per_task: int = 40
    hi_conf: float = 0.9
    control_share: float = 0.1
    repeat_share: float = 0.1
    """Share of items re-presented blind in a second sitting (intra-rater consistency)."""
    seed: int = 0


PLAN_SOURCES = (
    ("jev", "jev_main", "jev__canonical"),
    ("jev_noul", "jev_format", "jev__noul_format"),
    ("luna", "llm_main", "luna-verbal__canonical"),
    ("haiku", "llm_main", "haiku-verbal__canonical"),
    ("oss", "llm_main", "oss-verbal__canonical"),
    ("sonnet", "llm_frontier", "sonnet-verbal__canonical"),
)
"""Pool from docs/ANALYSIS_PLAN.md section 9."""
PLAN_TASKS = ("mfrc", "annomi_client", "annomi_therapist", "talkmoves_teacher", "talkmoves_student")


def units(task: Task, codebook: Codebook, preds: dict[str, pd.DataFrame], ids: set[str] | None = None) -> pd.DataFrame:
    """One row per unit with the human label, each coder's label/confidence (NaN where the coder did
    not code the unit), and disagreement flags over the coders that did."""
    covered = set.union(*(set(p.segment_id) for p in preds.values()))
    ids = covered if ids is None else covered & ids
    ref = reference(task.labels[task.labels.segment_id.isin(ids)], codebook)
    if codebook.mode == "exclusive":
        out = ref[["segment_id", "ref_code"]].rename(columns={"ref_code": "human"}).assign(code_id=None)
        for coder, p in preds.items():
            wide = p.pivot_table(index="segment_id", columns="code_id", values="p").reindex(columns=codebook.code_ids)
            ranked = wide.apply(lambda r: r.sort_values(ascending=False, kind="stable").index.tolist(), axis=1)
            out[f"{coder}_label"] = out.segment_id.map(ranked.str[0])
            out[f"{coder}_conf"] = out.segment_id.map(wide.max(axis=1))
            out[f"{coder}_runner_up"] = out.segment_id.map(ranked.str[1])
    else:
        out = ref[["segment_id", "code_id", "y"]].rename(columns={"y": "human"})
        for coder, p in preds.items():
            m = out.merge(p[["segment_id", "code_id", "p"]], on=["segment_id", "code_id"], how="left")
            out[f"{coder}_label"] = np.where(m.p.isna(), np.nan, (m.p >= 0.5).astype(float))
            out[f"{coder}_conf"] = np.maximum(m.p, 1 - m.p).to_numpy()
    disagree = np.column_stack([out[f"{c}_label"].notna() & (out[f"{c}_label"] != out.human) for c in preds])
    conf = np.column_stack([out[f"{c}_conf"].fillna(-np.inf) for c in preds])
    out["n_disagree"] = disagree.sum(axis=1)
    out["max_conf_disagree"] = np.where(disagree, conf, -np.inf).max(axis=1)
    return out.reset_index(drop=True)


def select(u: pd.DataFrame, cfg: ExportConfig, rng: np.random.Generator) -> pd.DataFrame:
    n_ctrl = int(round(cfg.n_per_task * cfg.control_share))
    n_dis = cfg.n_per_task - n_ctrl
    dis, agree = u[u.n_disagree > 0], u[u.n_disagree == 0]
    hi, rest = dis[dis.max_conf_disagree >= cfg.hi_conf], dis[dis.max_conf_disagree < cfg.hi_conf]

    def take(df: pd.DataFrame, k: int, stratum: str) -> pd.DataFrame:
        k = min(k, len(df))
        idx = rng.choice(len(df), size=k, replace=False) if k else []
        return df.iloc[idx].assign(stratum=stratum, inclusion=k / len(df) if len(df) else np.nan)

    a = take(hi, n_dis, "high_conf_disagreement")
    b = take(rest, n_dis - len(a), "other_disagreement")
    c = take(agree, n_ctrl, "control")
    return pd.concat([a, b, c], ignore_index=True)


def describe(code_id: str, codebook: Codebook) -> str:
    code = next(c for c in codebook.all_options if c.id == code_id)
    return f"{code.label}: {code.definition}"


def _label(v) -> str | None:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    return str(int(v)) if isinstance(v, (float, np.floating, int, np.integer)) else str(v)


def blind(sel: pd.DataFrame, task: Task, codebook: Codebook, coders: tuple[str, ...], rng: np.random.Generator):
    """Adjudicator-facing rows and the matching answer-key rows."""
    segs = task.segments.set_index("segment_id")
    sheet, key = [], []
    for r in sel.sample(frac=1, random_state=int(rng.integers(1 << 31))).itertuples():
        item = "".join(rng.choice(list(LETTERS + "23456789"), size=6))
        seg = segs.loc[r.segment_id]
        base = {
            "item": item,
            "task": task.name,
            "setting": seg.setting if isinstance(seg.setting, str) else "",
            "context": seg.context if isinstance(seg.context, str) else "",
            "segment": seg.text,
        }
        if codebook.mode == "exclusive":
            cands = {r.human} | {getattr(r, f"{c}_label") for c in coders if isinstance(getattr(r, f"{c}_label"), str)}
            if len(cands) == 1:  # control: pair the agreed label with the first coder's runner-up
                cands.add(getattr(r, f"{coders[0]}_runner_up"))
            order = [c for c in codebook.code_ids if c in cands]
            rng.shuffle(order)
            letters = {LETTERS[i]: c for i, c in enumerate(order)}
            question = f"{codebook.instructions} Which option(s) are correct under the codebook?"
            options = "\n".join(f"{L}) {describe(c, codebook)}" for L, c in letters.items())
        else:
            letters = {}
            question = f"Does the {codebook.unit} express the code \"{describe(r.code_id, codebook).split(':')[0]}\"? (Y / N / ?)"
            options = describe(r.code_id, codebook)
        sheet.append(base | {"question": question, "options": options, "answer": "", "note": ""})
        key.append(
            {
                "item": item,
                "task": task.name,
                "segment_id": r.segment_id,
                "code_id": r.code_id,
                "stratum": r.stratum,
                "inclusion": r.inclusion,
                "letters": json.dumps(letters),
                # Labels are strings throughout ("0"/"1" for multi codebooks) so tasks share one schema;
                # None where that coder did not code the unit.
                "human": _label(r.human),
                **{f"{c}_label": _label(getattr(r, f"{c}_label")) for c in coders},
                **{f"{c}_conf": float(getattr(r, f"{c}_conf")) for c in coders},
            }
        )
    return sheet, key


INSTRUCTIONS = """# Adjudication instructions

You are judging how qualitative codes should be applied, **strictly under the codebook as written**
(see `codebooks.md`), not under your own theory of the construct.

Each row is one item. Read `setting` and `context` (earlier turns, which are not themselves coded),
then the `segment`, then answer in the `answer` column:

- **Choice items** (options labelled A, B, ...): write the letter of the correct option. If more than
  one option is defensible under the codebook, write all of them joined with `+` (e.g. `A+B`). If none
  is right, write `none` and, if you can, name the right code in `note`.
- **Yes/no items** ("Does the ... express the code ...?"): write `Y`, `N`, or `?` (genuinely ambiguous
  under the codebook).

Options are shown in random order. Some items were chosen because coders disagreed; others were not. If there is a second workbook
(`adjudication_part2_later.xlsx`), do it in a separate sitting at least a day later, without looking
back at the first one.
Nothing about an option's position or wording tells you where it came from. Use `note` for anything
that would help (a codebook rule you relied on, a missing rule, a transcription problem).
"""


def codebooks_markdown(codebooks: list[Codebook]) -> str:
    parts = ["# Codebooks\n"]
    for cb in codebooks:
        parts.append(f"## {cb.name} ({cb.unit}; {'one code' if cb.mode == 'exclusive' else 'any number of codes'})\n")
        parts.append(f"{cb.instructions}\n\n_Source: {cb.source}_\n")
        for c in cb.all_options:
            parts.append(f"### {c.label}\n\n{c.definition}\n")
            for title, items in (("Applies when", c.include), ("Does not apply when", c.exclude),
                                 ("Examples", c.examples_pos), ("Non-examples", c.examples_neg)):
                if items:
                    parts.append(f"**{title}:**\n" + "\n".join(f"- {i}" for i in items) + "\n")
    return "\n".join(parts)


def write_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    widths = {"item": 9, "task": 16, "setting": 28, "context": 48, "segment": 48, "question": 40,
              "options": 64, "answer": 10, "note": 32}
    wb = Workbook()
    ws = wb.active
    ws.title = "items"
    cols = list(widths)
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    yn = DataValidation(type="list", formula1='"Y,N,?"', allow_blank=True)
    ws.add_data_validation(yn)
    for i, row in enumerate(rows, start=2):
        ws.append([row[c] for c in cols])
        if row["question"].endswith("(Y / N / ?)"):
            yn.add(f"H{i}")
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths[c]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def export(cfg: ExportConfig, batch: str) -> Path:
    from margin import sampling

    rng = np.random.default_rng(cfg.seed)
    out = ADJUDICATION / batch
    handout = out / "to_adjudicator"
    handout.mkdir(parents=True, exist_ok=True)
    names = tuple(n for n, _, _ in cfg.sources)
    sheet, key, codebooks, counts = [], [], [], {}
    for task_name in cfg.tasks:
        codebook = Codebook.load(f"codebooks/{task_name}.yaml")
        task = Task.load(task_name)
        preds = {n: pd.read_parquet(RUNS / exp / task_name / f"{stem}.parquet") for n, exp, stem in cfg.sources
                 if (RUNS / exp / task_name / f"{stem}.parquet").exists()}
        sample = sampling.load(task_name, "main")
        ids = set(sample.segment_id if cfg.split is None else sample[sample.split == cfg.split].segment_id)
        u = units(task, codebook, preds, ids)
        for n in names:  # sources missing for this task get empty columns so the key schema is stable
            if f"{n}_label" not in u:
                u[f"{n}_label"], u[f"{n}_conf"] = np.nan, np.nan
        sel = select(u, cfg, rng)
        s_, k_ = blind(sel, task, codebook, names, rng)
        sheet += s_
        key += k_
        codebooks.append(codebook)
        counts[task_name] = {"units": len(u), "disagreements": int((u.n_disagree > 0).sum()),
                             "high_conf_disagreements": int((u.max_conf_disagree >= cfg.hi_conf).sum()),
                             "selected": sel.stratum.value_counts().to_dict(), "sources": sorted(preds)}
    # Second sitting: a blind repeat of a random share of items under new ids, in its own workbook.
    n_rep = int(round(cfg.repeat_share * len(sheet)))
    picks = rng.choice(len(sheet), size=n_rep, replace=False) if n_rep else []
    part2, used = [], {r["item"] for r in sheet}
    for i in picks:
        item = "".join(rng.choice(list(LETTERS + "23456789"), size=6))
        while item in used:
            item = "".join(rng.choice(list(LETTERS + "23456789"), size=6))
        used.add(item)
        part2.append(sheet[i] | {"item": item})
        key.append(key[i] | {"item": item, "stratum": "repeat", "repeat_of": key[i]["item"]})
    rng.shuffle(part2)
    write_xlsx(sheet, handout / "adjudication.xlsx")
    if part2:
        write_xlsx(part2, handout / "adjudication_part2_later.xlsx")
    (handout / "codebooks.md").write_text(codebooks_markdown(codebooks))
    (handout / "INSTRUCTIONS.md").write_text(INSTRUCTIONS)
    pd.DataFrame(key).to_parquet(out / "key.parquet", index=False)
    (out / "manifest.json").write_text(json.dumps({"config": cfg.__dict__, "counts": counts, "repeats": n_rep}, indent=2, default=str))
    return out


def read_answers(path: Path) -> pd.DataFrame:
    if path.suffix == ".xlsx":
        return pd.read_excel(path, sheet_name="items", dtype=str).fillna("")
    return pd.read_csv(path, dtype=str).fillna("")


def parse(answer: str, letters: dict[str, str], multi: bool) -> tuple[list, bool]:
    """-> (acceptable labels, ambiguous). Raises ValueError on an unreadable answer."""
    a = answer.strip().lower()
    if multi:
        if a not in MULTI_ANSWERS:
            raise ValueError(f"expected Y/N/?, got {answer!r}")
        v = MULTI_ANSWERS[a]
        return ([] if v is None else [str(v)]), v is None
    if a in NONE_ANSWERS:
        return [], False
    picked = [p.strip().upper() for p in re.split(r"[+,/\s]+", a) if p.strip()]
    if not picked or any(p not in letters for p in picked):
        raise ValueError(f"expected letters from {sorted(letters)} (joined with +) or 'none', got {answer!r}")
    return [letters[p] for p in picked], False


def import_answers(batch: str, paths: list[Path], adjudicator: str) -> pd.DataFrame:
    """Read one or more filled workbooks (main + second sitting), unblind against the key, save."""
    out = ADJUDICATION / batch
    key = pd.read_parquet(out / "key.parquet")
    answers = pd.concat([read_answers(Path(p))[["item", "answer", "note"]] for p in paths], ignore_index=True)
    d = key.merge(answers, on="item", how="inner").fillna({"answer": "", "note": ""})
    rows, problems = [], []
    for r in d.itertuples():
        if not r.answer.strip():
            continue
        try:
            acceptable, ambiguous = parse(r.answer, json.loads(r.letters), multi=r.code_id is not None and pd.notna(r.code_id))
        except ValueError as err:
            problems.append(f"{r.item}: {err}")
            continue
        rows.append({"item": r.item, "acceptable": acceptable, "ambiguous": ambiguous, "note": r.note})
    if problems:
        raise ValueError("unreadable answers:\n" + "\n".join(problems))
    verdicts = key.merge(pd.DataFrame(rows), on="item").assign(adjudicator=adjudicator)
    verdicts.to_parquet(out / f"verdicts_{adjudicator}.parquet", index=False)
    return verdicts


def _verdict(acceptable, ambiguous) -> str:
    return "?" if ambiguous else json.dumps(sorted(list(acceptable), key=str))


def intra_rater(verdicts: pd.DataFrame) -> dict:
    """Agreement between each repeated item and its original (same adjudicator, blind second sitting)."""
    if "repeat_of" not in verdicts or verdicts.repeat_of.isna().all():
        return {"repeats": 0}
    v = verdicts.assign(verdict=[_verdict(a, m) for a, m in zip(verdicts.acceptable, verdicts.ambiguous)])
    orig = v.set_index("item").verdict
    rep = v[v.repeat_of.notna() & v.repeat_of.isin(orig.index)]
    same = (rep.verdict.to_numpy() == orig.loc[rep.repeat_of].to_numpy())
    return {"repeats": len(rep), "self_agreement": float(same.mean()) if len(rep) else float("nan")}


def summary(verdicts: pd.DataFrame, coders: list[str]) -> pd.DataFrame:
    """Per task x stratum: how often each source's label was judged acceptable (over the items that
    source coded). Repeats are excluded here; see intra_rater()."""
    v = verdicts[~verdicts.ambiguous & (verdicts.stratum != "repeat")].copy()

    def ok(col: str) -> pd.Series:
        # A source that did not code an item has no label there (None or NaN once read back from parquet).
        return pd.Series([np.nan if pd.isna(lab) else float(lab in acc) for lab, acc in zip(v[col], v.acceptable)], index=v.index)

    v["human_ok"] = ok("human")
    for c in coders:
        v[f"{c}_ok"] = ok(f"{c}_label")
    v["none_ok"] = (v.acceptable.map(len) == 0).astype(float)
    cols = ["human_ok"] + [f"{c}_ok" for c in coders] + ["none_ok"]
    table = v.groupby(["task", "stratum"])[cols].mean().round(3)
    table.insert(0, "n", v.groupby(["task", "stratum"]).size())
    table["ambiguous"] = verdicts[verdicts.stratum != "repeat"].groupby(["task", "stratum"]).ambiguous.sum()
    return table.reset_index()


def load_verdicts(batch: str) -> pd.DataFrame:
    """Verdicts for the adjudicated reference: repeats excluded; where several adjudicators judged an
    item, it is kept only if they agree."""
    final = ADJUDICATION / batch / "verdicts_final.parquet"
    if final.exists():  # panel + author protocol (deviation 2): already resolved, one verdict per item
        v = pd.read_parquet(final)
        v["acceptable"] = v.acceptable.map(lambda a: sorted(a.tolist() if hasattr(a, "tolist") else a, key=str))
        return v
    files = sorted((ADJUDICATION / batch).glob("verdicts_*.parquet"))
    if not files:
        raise FileNotFoundError(f"no verdicts in {ADJUDICATION / batch}")
    v = pd.concat([pd.read_parquet(f) for f in files], ignore_index=True)
    v = v[v.stratum != "repeat"].copy()
    v["acceptable"] = v.acceptable.map(lambda a: sorted(a.tolist() if hasattr(a, "tolist") else a, key=str))
    v["verdict"] = [_verdict(a, m) for a, m in zip(v.acceptable, v.ambiguous)]
    agree = v.groupby("item").verdict.nunique() == 1
    return v[v.item.isin(agree[agree].index)].drop_duplicates("item")
